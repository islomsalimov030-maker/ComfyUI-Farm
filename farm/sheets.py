"""Таблица сцен (WS-03): чтение Google Sheets, проверка строк, синхронизация в БД, статусы, выгрузка финалов.

Схема листов — farm/sheet_schema.yaml (колонки ищутся по заголовку). Контракт — CONTRACTS §10 (v0.4).

Точки входа:
  sync(db)                         строки «В работу» → shots (DB.add_shot), ошибки/статусы → колонка «Статус»
  on_shot_update(shot, links, scores)   воркер сообщает о шоте → Статус/Ссылки/Оценки/Обновлено
  upload_final(asset_path, shot)   финал → Google Drive (сверка md5 Drive) → ссылка | None
  python -m farm.sheets check|sync [--xlsx файл]   проверка без БД / синхронизация из консоли

Настройки — секция `sheets:` в $FARM_HOME/farm.yaml (см. DEFAULTS ниже и docs/SHEETS_SETUP.md).
Нет настроек или ключа — модуль молчит (заглушка), воркер не падает.
"""
from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import logging
import re
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional

from .config import deep_merge, farm_home, load_config
from .library import CardError, Library, cache_path, load_schema, parse_anchors, parse_slots

log = logging.getLogger("farm.sheets")

DEFAULTS = {
    "enabled": True,
    "spreadsheet_id": None,                 # из адреса таблицы: docs.google.com/spreadsheets/d/<ID>/edit
    "service_account": "secrets/google_sa.json",   # от FARM_HOME
    "xlsx": None,                           # офлайн: путь к выгрузке .xlsx вместо Google API
    "check_files": "auto",                  # auto — проверять файлы refs/, если папка есть (на поде)
    "timezone": "Asia/Tashkent",
    "flush_s": 3.0,                         # запись статусов от воркера пачками
}
DRIVE_DEFAULTS = {
    "folder_id": None,                      # папка финалов на Google Drive
    "oauth_token": "secrets/google_oauth_token.json",  # токен пользователя (у SA нет квоты на «Мой диск»)
}

SCHEMA = load_schema()
S_COLS = SCHEMA["scenes"]["columns"]
LIM = SCHEMA["limits"]
LISTS = SCHEMA["lists"]
STATUS = SCHEMA["statuses"]
ORCH_KEYS = [c["key"] for c in S_COLS if c.get("orchestrator")]
# поля, правка которых после старта не применяется (меняют граф/вид работы)
STRUCT_KEYS = ("kind", "start_from", "end_with", "depends", "profile")
ACTIVE_STATES = {"frame_gen", "frame_qa", "frame_retry", "video_gen", "video_qa", "video_retry", "post_fix"}

_CYR = re.compile(r"[А-Яа-яЁё]")
_DRIVE_ID = re.compile(r"(?:/d/|[?&]id=)([A-Za-z0-9_-]{20,})")


class SheetError(RuntimeError):
    """Таблица не читается как надо (нет листа/колонки). Текст — по-русски."""


# ======================================================================== бэкенды
class MemoryBackend:
    """Таблица в памяти: {лист: [[ячейки строки], ...]} (строка 1 — заголовок). Для тестов."""

    def __init__(self, tables: dict[str, list[list[Any]]]):
        self.tables = {k: [list(r) for r in v] for k, v in tables.items()}
        self.writes: list[tuple[str, int, int, Any]] = []

    def read(self, sheet: str) -> list[list[Any]]:
        if sheet not in self.tables:
            raise SheetError(f"нет листа «{sheet}»")
        return [list(r) for r in self.tables[sheet]]

    def write(self, sheet: str, cells: list[tuple[int, int, Any]]) -> None:
        t = self.tables[sheet]
        for r, c, v in cells:                       # 1-based
            while len(t) < r:
                t.append([])
            row = t[r - 1]
            while len(row) < c:
                row.append("")
            row[c - 1] = v
            self.writes.append((sheet, r, c, v))


class XlsxBackend:
    """Офлайн-режим: выгрузка таблицы в .xlsx (Файл → Скачать). Запись — в тот же файл."""

    def __init__(self, path: str | Path):
        self.path = Path(path)

    def _wb(self):
        import openpyxl  # лениво: на поде не обязателен
        return openpyxl.load_workbook(self.path)

    def read(self, sheet: str) -> list[list[Any]]:
        wb = self._wb()
        if sheet not in wb.sheetnames:
            raise SheetError(f"нет листа «{sheet}» в {self.path.name}")
        return [["" if v is None else v for v in row] for row in wb[sheet].iter_rows(values_only=True)]

    def write(self, sheet: str, cells: list[tuple[int, int, Any]]) -> None:
        wb = self._wb()
        ws = wb[sheet]
        for r, c, v in cells:
            ws.cell(row=r, column=c, value=v)
        wb.save(self.path)


class GSheetsBackend:
    """Google Sheets через сервисный аккаунт (gspread, MIT)."""

    def __init__(self, spreadsheet_id: str, sa_json: str | Path):
        import gspread  # лениво: модуль импортируется и без gspread (заглушка)
        self._gspread = gspread
        self.gc = gspread.service_account(filename=str(sa_json))
        self.sh = self.gc.open_by_key(spreadsheet_id)
        self._ws: dict[str, Any] = {}

    def _sheet(self, name: str):
        if name not in self._ws:
            try:
                self._ws[name] = self.sh.worksheet(name)
            except self._gspread.WorksheetNotFound:
                raise SheetError(f"нет листа «{name}»") from None
        return self._ws[name]

    def _retry(self, fn, *a, **kw):
        delay = 2.0
        for i in range(6):
            try:
                return fn(*a, **kw)
            except self._gspread.exceptions.APIError as e:
                code = getattr(getattr(e, "response", None), "status_code", 0)
                if code in (429, 500, 502, 503) and i < 5:
                    time.sleep(delay)
                    delay = min(delay * 2, 60)
                    continue
                raise

    def read(self, sheet: str) -> list[list[Any]]:
        return self._retry(self._sheet(sheet).get_all_values)

    def write(self, sheet: str, cells: list[tuple[int, int, Any]]) -> None:
        if not cells:
            return
        from gspread.utils import rowcol_to_a1
        data = [{"range": rowcol_to_a1(r, c), "values": [[v]]} for r, c, v in cells]
        self._retry(self._sheet(sheet).batch_update, data, value_input_option="RAW")


# ======================================================================== чтение листа
def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace("*", "").strip().lower())


def header_map(header: list[Any], columns: list[dict], sheet: str) -> dict[str, int]:
    """{key: индекс колонки 0-based}. Нет обязательной колонки → SheetError.

    Совпадение по заголовку целиком, иначе по части до запятой («Свет, тон» ↔ «Свет»): листы,
    созданные до схемы v1 с кавычками, имеют обрезанные заголовки; одинаковые берутся по порядку.
    """
    pos: dict[str, list[int]] = {}
    for i, h in enumerate(header):
        if _norm_header(h):
            pos.setdefault(_norm_header(h), []).append(i)
    used: set[int] = set()

    def take(name: str) -> Optional[int]:
        return next((i for i in pos.get(name, []) if i not in used), None)

    out = {}
    for c in columns:
        full = _norm_header(c["header"])
        i = take(full)
        if i is None and "," in full:
            i = take(full.split(",")[0].strip())
        if i is not None:
            used.add(i)
            out[c["key"]] = i
        elif c.get("required") or c.get("orchestrator"):
            raise SheetError(f"в листе «{sheet}» нет колонки «{c['header']}»")
    return out


def read_rows(values: list[list[Any]], columns: list[dict], sheet: str) -> list[tuple[int, dict]]:
    """[(номер строки в листе, {key: текст})], пустые строки пропускаются."""
    if not values:
        raise SheetError(f"лист «{sheet}» пуст (нет строки заголовков)")
    hm = header_map(values[0], columns, sheet)
    boxes = {c["key"] for c in columns if c.get("type") == "checkbox"}
    out = []
    for n, row in enumerate(values[1:], start=2):
        d = {k: (row[i] if i < len(row) else "") for k, i in hm.items()}
        d = {k: ("" if v is None else (v if isinstance(v, bool) else str(v).strip())) for k, v in d.items()}
        if any(not _blank(v, k in boxes) for k, v in d.items() if not _is_orch(k, columns)):
            out.append((n, d))
    return out


def _blank(v: Any, checkbox: bool) -> bool:
    """Пустая ячейка; невыставленная галочка (FALSE) строку не делает непустой."""
    if v in ("", None, False):
        return True
    return checkbox and str(v).strip().lower() in ("false", "ложь")


def _is_orch(key: str, columns: list[dict]) -> bool:
    return any(c["key"] == key and c.get("orchestrator") for c in columns)


# ======================================================================== нормализация
_TRUE = {"да", "yes", "y", "true", "1", "✓", "✔", "истина", "+"}
_FALSE = {"нет", "no", "n", "false", "0", "ложь", "-", "—"}


def parse_bool(v: Any) -> Optional[bool]:
    if isinstance(v, bool):
        return v
    s = str(v or "").strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    return None


def parse_number(v: Any) -> Optional[float]:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v or "").strip().lower()
    m = re.fullmatch(r"~?\s*(\d+(?:[.,]\d+)?)\s*(?:с|сек|s|sec)?\.?", s)
    return float(m.group(1).replace(",", ".")) if m else None


def split_ids(v: Any) -> list[str]:
    return [t for t in re.split(r"[,;\n\s]+", str(v or "").strip()) if t]


def is_drive_link(s: str) -> bool:
    return "drive.google.com" in s or "docs.google.com" in s


def drive_file_id(s: str) -> Optional[str]:
    m = _DRIVE_ID.search(s or "")
    return m.group(1) if m else None


# ======================================================================== проверка строк
@dataclass
class Issue:
    code: str
    key: str            # колонка (ключ схемы) или ""
    msg: str
    severity: str = "error"   # error | warn


@dataclass
class RowCheck:
    row: int
    raw: dict
    norm: dict = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)

    @property
    def id(self) -> str:
        return self.norm.get("id") or str(self.raw.get("id") or "")

    @property
    def ready(self) -> bool:
        return bool(self.norm.get("ready"))

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.severity == "error"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.severity == "warn"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def err(self, code, key, msg):
        self.issues.append(Issue(code, key, msg, "error"))

    def warn(self, code, key, msg):
        self.issues.append(Issue(code, key, msg, "warn"))

    def status_text(self) -> str:
        hdr = {c["key"]: c["header"] for c in S_COLS}
        parts = [f"{hdr.get(i.key, '')}: {i.msg}" if i.key else i.msg for i in self.errors]
        return (STATUS["invalid"] + "; ".join(parts))[:900]

    def warn_text(self) -> str:
        return "; ".join(i.msg for i in self.warnings)[:400]


def _suggest(value: str, choices: Iterable[str]) -> str:
    m = difflib.get_close_matches(value, list(choices), n=1, cutoff=0.6)
    return f" — может, «{m[0]}»?" if m else ""


def _check_id(rc: RowCheck, key: str, value: str, pattern: str, what: str) -> bool:
    if _CYR.search(value):
        rc.err("E_ID_CYR", key, f"в {what} «{value}» русские буквы — нужна латиница (частая ошибка: С, Е, А, О, Р)")
        return False
    if not re.fullmatch(pattern, value):
        rc.err("E_ID_FORMAT", key, f"{what} «{value}»: только латиница, цифры, дефис")
        return False
    return True


def _ref(rc: RowCheck, key: str, lib: Library, libname: str, value: str, what: str) -> Optional[str]:
    v = value.strip().lower()
    if _CYR.search(v):
        rc.err("E_ID_CYR", key, f"«{value}»: русские буквы в ID {what}")
        return None
    if v not in lib.ids(libname):
        rc.err("E_REF_UNKNOWN", key, f"нет {what} «{value}» в справочнике{_suggest(v, lib.ids(libname))}")
        return None
    return v


def validate_row(raw: dict, row: int, lib: Library, active_profiles: Optional[set] = None) -> RowCheck:
    """Проверка одной строки без учёта соседей (дубли, связи, циклы — в validate)."""
    rc = RowCheck(row=row, raw=raw)
    n = rc.norm
    g = lambda k: raw.get(k, "")            # noqa: E731

    # ID
    sid = str(g("id")).strip()
    if not sid:
        rc.err("E_ID_EMPTY", "id", "нет ID")
    elif _check_id(rc, "id", sid, LIM["shot_id_re"], "ID"):
        n["id"] = sid.upper()
    n["name"] = str(g("name"))
    rdy = parse_bool(g("ready"))
    n["ready"] = bool(rdy)
    n["batch"] = str(g("batch"))

    # приоритет
    p = g("priority")
    if p in ("", None):
        n["priority"] = LIM["priority_default"]
    else:
        pv = parse_number(p)
        lo, hi = LIM["priority"]
        if pv is None or pv != int(pv) or not lo <= pv <= hi:
            rc.err("E_PRIORITY", "priority", f"«{p}»: нужно целое от {lo} до {hi}")
        else:
            n["priority"] = int(pv)

    # профиль
    prof = str(g("profile")).strip().lower()
    if not prof:
        rc.err("E_REQUIRED", "profile", "заполните")
    else:
        known = lib.ids("profiles")
        if prof not in known:
            rc.err("E_PROFILE_UNKNOWN", "profile", f"нет профиля «{prof}» на листе «Профили»{_suggest(prof, known)}")
        elif active_profiles is not None and prof not in active_profiles:
            rc.err("E_PROFILE_OFF", "profile", f"профиль «{prof}» выключен (лист «Профили», «Активен»)")
        else:
            n["profile"] = prof

    # тип
    k = str(g("kind")).strip().lower()
    kinds = {**LISTS["kind"], "video": "video", "frame": "frame", "фото": "frame", "картинка": "frame"}
    if not k:
        n["kind"] = "video"                 # v2: колонки «Тип» может не быть — по умолчанию видео
    elif k not in kinds:
        rc.err("E_KIND", "kind", f"«{k}»: нужно «видео» или «кадр»")
    else:
        n["kind"] = kinds[k]
    video = n.get("kind") == "video"

    # справочники
    ch = str(g("character")).strip()
    if not ch:
        rc.err("E_REQUIRED", "character", "заполните")
    else:
        n["character"] = _ref(rc, "character", lib, "characters", ch, "персонажа")

    pl = str(g("plate")).strip().lower().replace(" ", "")
    zn = str(g("zone")).strip().lower().replace(" ", "")
    n["plate"] = n["zone"] = None
    if zn:                                   # v2: зона (location sheet) вместо плиты
        n["zone"] = _ref(rc, "zone", lib, "zones", zn, "зоны")
        n["location"] = zn.split("_")[0]
    if not pl:
        if not zn:
            rc.err("E_REQUIRED", "zone", "заполните «Зона» (или «Локация / плита» для таблицы v1)")
    elif not re.fullmatch(LIM["plate_id_re"], pl):
        rc.err("E_PLATE_FORMAT", "plate", f"«{g('plate')}»: нужен вид локация/зона латиницей, например apartment/sofa")
    else:
        n["plate"] = _ref(rc, "plate", lib, "plates", pl, "плиты")
        n["location"] = pl.split("/")[0]

    of = str(g("outfit")).strip()
    if of:
        n["outfit"] = _ref(rc, "outfit", lib, "outfits", of, "наряда")
    else:
        n["outfit"] = None
        if not n.get("zone"):                # v2: пустой наряд = базовый наряд, это не ошибка
            rc.warn("W_OUTFIT_EMPTY", "outfit", "наряд не указан")

    objs = []
    for o in split_ids(g("objects")):
        v = _ref(rc, "objects", lib, "objects", o, "объекта")
        if v:
            objs.append(v)
    n["objects"] = objs

    # v2: класс позы, план, лицо в кадре
    pc = str(g("pose_class")).strip().lower()
    n["pose_class"] = LISTS.get("pose_class", {}).get(pc) if pc else None
    if pc and n["pose_class"] is None:
        rc.err("E_POSE_CLASS", "pose_class", f"«{pc}»: нужно стоя / сидя / присед / лежа")
    pln = str(g("plan")).strip().lower()
    n["plan"] = LISTS.get("plan", {}).get(pln) if pln else None
    if pln and n["plan"] is None:
        rc.warn("W_PLAN", "plan", f"«{pln}»: ожидается общий / средний / крупный — будет авто")
    n["face_in_frame"] = str(g("face_in_frame")).strip().lower() or None

    # тексты
    for key in ("framing", "pose", "action", "focus", "amplitude", "base", "emotion", "note"):
        n[key] = str(g(key))
        if len(n[key]) > LIM["text_max"]:
            rc.warn("W_TEXT_LONG", key, f"текст длиннее {LIM['text_max']} знаков — сократите")
    if not n["pose"]:
        rc.err("E_REQUIRED", "pose", "опишите стартовую позу")
    if video and not n["action"]:
        rc.err("E_ACTION_REQUIRED", "action", "для видео опишите действие")
    n["mood"] = str(g("mood")).strip().lower()

    # петля, длительность
    lp = g("loop")
    lb = parse_bool(lp)
    if video:
        if lp in ("", None):
            rc.err("E_LOOP", "loop", "для видео укажите да или нет")
        elif lb is None:
            rc.err("E_LOOP", "loop", f"«{lp}»: нужно да или нет")
    n["loop"] = lb

    d = g("duration")
    dv = parse_number(d) if d not in ("", None) else None
    lo, hi = LIM["duration_s"]
    if video:
        if d in ("", None):
            rc.err("E_DURATION", "duration", f"для видео укажите секунды ({lo}–{hi})")
        elif dv is None:
            rc.err("E_DURATION", "duration", f"«{d}» — не число секунд")
        elif not lo <= dv <= hi:
            rc.err("E_DURATION", "duration", f"{d} с: движок берёт от {lo} до {hi} с")
        else:
            step = LIM["duration_step"]
            r = round(dv / step) * step
            if abs(r - dv) > 1e-9:
                rc.warn("W_DURATION_STEP", "duration", f"длительность округлена до {r:g} с")
            n["duration"] = r
    elif d not in ("", None):
        rc.warn("W_DURATION_FRAME", "duration", "для кадра длительность не нужна")
        n["duration"] = None
    else:
        n["duration"] = None

    # связи между шотами (существование — в validate)
    for key in ("start_from", "end_with"):
        v = str(g(key)).strip()
        n[key] = None
        if v:
            if _check_id(rc, key, v, LIM["shot_id_re"], "ID шота"):
                n[key] = v.upper()
                if not video:
                    rc.err("E_LINK_KIND", key, "только для типа «видео»")
    n["depends"] = []
    for v in split_ids(g("depends")):
        if _check_id(rc, "depends", v, LIM["shot_id_re"], "ID шота"):
            n["depends"].append(v.upper())
    if n["end_with"] and lb:
        rc.err("E_LOOP_END", "loop", "с «Закончить кадром шота» петля невозможна — поставьте «нет»")

    # референсы
    for key in ("pose_ref", "motion_ref", "location_image"):
        v = str(g(key)).strip()
        n[key] = v or None
        if v and v.startswith("http") and not (is_drive_link(v) and drive_file_id(v)):
            rc.err("E_LINK", key, "нужна ссылка на файл Google Drive (…/file/d/<id>/…) или путь в refs/")
    if n["motion_ref"] and not video:
        rc.warn("W_MOTION_FRAME", "motion_ref", "для кадра видео-реф не используется")
    return rc


def validate(rows: list[tuple[int, dict]], lib: Library, active_profiles: Optional[set] = None,
             check_files: bool = False) -> list[RowCheck]:
    """Проверка всех строк листа «Сцены» (с учётом дублей, связей и циклов)."""
    checks = [validate_row(raw, n, lib, active_profiles) for n, raw in rows]
    by_id: dict[str, list[RowCheck]] = {}
    for rc in checks:
        if rc.norm.get("id"):
            by_id.setdefault(rc.norm["id"], []).append(rc)
    for sid, group in by_id.items():
        if len(group) > 1:
            rows_s = ", ".join(str(r.row) for r in group)
            for rc in group:
                rc.err("E_ID_DUP", "id", f"ID «{sid}» повторяется (строки {rows_s})")

    hdr = {c["key"]: c["header"] for c in S_COLS}
    for rc in checks:
        n = rc.norm
        for key in ("start_from", "end_with", "depends"):
            targets = n.get(key) or []
            targets = [targets] if isinstance(targets, str) else targets
            for t in targets:
                if t == n.get("id"):
                    rc.err("E_LINK_SELF", key, "шот не может ссылаться сам на себя")
                elif t not in by_id:
                    rc.err("E_LINK_UNKNOWN", key, f"нет шота «{t}» в листе{_suggest(t, by_id)}")
                else:
                    src = by_id[t][0]
                    if rc.ready and not (src.ready and src.ok):
                        rc.err("E_LINK_NOT_READY", key, f"шот «{t}» не «В работу» или с ошибками")
                    if key == "start_from" and src.norm.get("kind") == "frame":
                        rc.err("E_LINK_KIND", key, f"«{t}» — кадр без видео, начать с его последнего кадра нельзя")
        if check_files and rc.ok:
            _check_files(rc, lib)
    # циклы
    graph = {rc.norm["id"]: set(_deps(rc.norm)) for rc in checks if rc.norm.get("id")}
    for cyc in _cycles(graph):
        chain = " → ".join(cyc + [cyc[0]])
        for sid in cyc:
            for rc in by_id.get(sid, []):
                rc.err("E_CYCLE", "depends", f"замкнутая цепочка зависимостей: {chain}")
    return checks


def _deps(n: dict) -> list[str]:
    out = list(n.get("depends") or [])
    for k in ("start_from", "end_with"):
        if n.get(k) and n[k] not in out:
            out.append(n[k])
    return out


def _cycles(graph: dict[str, set]) -> list[list[str]]:
    seen, stack, on, out = set(), [], set(), []

    def dfs(v):
        seen.add(v)
        stack.append(v)
        on.add(v)
        for w in sorted(graph.get(v, ())):
            if w not in graph:
                continue
            if w in on:
                cyc = stack[stack.index(w):]
                if sorted(cyc) not in [sorted(c) for c in out]:
                    out.append(list(cyc))
            elif w not in seen:
                dfs(w)
        stack.pop()
        on.discard(v)

    for v in sorted(graph):
        if v not in seen:
            dfs(v)
    return out


def _check_files(rc: RowCheck, lib: Library) -> None:
    n = rc.norm
    if n.get("zone"):                        # v2: лист ракурсов зоны + эталон лица для QA видео; моды — по папкам
        need = [("zone", "zones", n["zone"], "sheet", "листа ракурсов зоны (tools/farm_prepare_refs.py)")]
        if n.get("kind") == "video":
            need.append(("character", "characters", n.get("character"), "face", "эталона лица"))
        for key, libname, id_, fkey, what in need:
            if id_ and lib.file(libname, id_, fkey) is None:
                rc.err("E_FILE_MISSING", key, f"нет файла {what}: {lib.expected_file(libname, id_, fkey)}")
        return
    need = [("plate", "plates", n.get("plate"), "plate", "плиты")]
    if n.get("kind") == "video" and not n.get("start_from"):
        need.append(("character", "characters", n.get("character"), "sheet", "листа персонажа"))
    elif n.get("kind") == "frame":
        need.append(("character", "characters", n.get("character"), "sheet", "листа персонажа"))
    if n.get("kind") == "video":
        need.append(("character", "characters", n.get("character"), "face", "эталона лица"))
    if n.get("outfit"):
        need.append(("outfit", "outfits", n["outfit"], "flatlay", "раскладки наряда"))
    for key, libname, id_, fkey, what in need:
        if id_ and lib.file(libname, id_, fkey) is None:
            rc.err("E_FILE_MISSING", key, f"нет файла {what}: {lib.expected_file(libname, id_, fkey)}")
    for key in ("pose_ref", "motion_ref", "location_image"):
        v = n.get(key)
        if v and not v.startswith("http") and _local_path(lib.home, v) is None:
            rc.err("E_FILE_MISSING", key, f"нет файла {v}")


def _local_path(home: Path, v: str) -> Optional[Path]:
    p = Path(v)
    for c in ([p] if p.is_absolute() else [home / v, home / "refs" / v]):
        if c.is_file():
            return c
    return None


# ======================================================================== строка → шот
def row_hash(n: dict) -> str:
    keys = ("profile", "kind", "character", "plate", "zone", "pose_class", "plan", "face_in_frame",
            "outfit", "objects", "framing", "pose", "action",
            "focus", "amplitude", "base", "emotion", "loop", "duration", "start_from", "end_with",
            "depends", "pose_ref", "motion_ref", "location_image")
    blob = json.dumps({k: n.get(k) for k in keys}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(blob.encode()).hexdigest()[:16]


def to_shot(rc: RowCheck, pose_ref: Optional[str] = None, motion_ref: Optional[str] = None,
            location_image: Optional[str] = None) -> dict:
    """RowCheck (ok) → dict для DB.add_shot (CONTRACTS §3 + extra_json v0.4)."""
    n = rc.norm
    lo, hi = LIM["priority"]
    return {
        "shot_id": n["id"], "sheet_row": rc.row, "profile": n["profile"], "kind": n["kind"],
        "character": n["character"], "location": n.get("location"), "plate": n.get("plate") or n.get("zone"),
        "outfit": n.get("outfit"), "objects_json": n.get("objects") or [],
        "pose_ru": n["pose"], "action_ru": n.get("action") or None,
        "loop": (int(bool(n["loop"])) if n.get("loop") is not None else None),
        "duration_s": n.get("duration"),
        "pose_ref": pose_ref, "motion_ref": motion_ref,
        "priority": hi + lo - n["priority"],         # в БД больше = раньше; в таблице 1 = срочно
        "depends_on_json": _deps(n),
        "extra_json": {
            "name": n.get("name") or None, "batch": n.get("batch") or None,
            "framing_ru": n.get("framing") or None, "focus_ru": n.get("focus") or None,
            "amplitude_ru": n.get("amplitude") or None, "base_ru": n.get("base") or None,
            "emotion_ru": n.get("emotion") or None, "mood": n.get("mood") or None,
            "start_from": n.get("start_from"), "end_with": n.get("end_with"),
            "zone": n.get("zone"), "pose_class": n.get("pose_class"), "plan": n.get("plan"),
            "face_in_frame": n.get("face_in_frame"),
            "location_image": location_image,   # 23.10: готовая картинка локации/листа строки → <Picture 1> как есть
            "row_hash": row_hash(n),
        },
    }


# ======================================================================== статус шота
def shot_status(shot: dict, waiting_for: Optional[list[str]] = None) -> str:
    st = shot.get("state") or "pending"
    if st == "pending" and waiting_for:
        return STATUS["waiting"] + ", ".join(waiting_for)
    text = STATUS.get(st, st)
    if st in ("frame_retry", "video_retry") and shot.get("stage_try"):
        text += f" {shot['stage_try']}"
    if st == "needs_human" and shot.get("last_error"):
        text += ": " + str(shot["last_error"])[:200]
    return text


def format_scores(scores: dict) -> str:
    names = {"frame": "кадр", "video": "видео"}
    parts = []
    for k, v in (scores or {}).items():
        if isinstance(v, float):
            v = f"{v:.2f}"
        parts.append(f"{names.get(k, k)} {v}")
    return " · ".join(parts)


# ======================================================================== основной класс
class Sheets:
    def __init__(self, backend, cfg: Optional[dict] = None, home: Optional[Path] = None,
                 creds_factory=None):
        self.backend = backend
        self.cfg = deep_merge(DEFAULTS, (cfg or {}).get("sheets"))
        self.drive_cfg = deep_merge(DRIVE_DEFAULTS, (cfg or {}).get("drive"))
        self.home = Path(home) if home else farm_home()
        self._creds_factory = creds_factory
        self._rowmap: dict[str, int] = {}
        self._cols: dict[str, int] = {}
        self._pending: dict[str, dict] = {}
        self._plock = threading.Lock()
        self._flusher: Optional[threading.Thread] = None

    # --- справочники -----------------------------------------------------------------
    def read_library(self) -> Library:
        data = {}
        for lib, spec in SCHEMA["libraries"].items():
            try:
                values = self.backend.read(spec["sheet"])
            except SheetError:
                if not spec.get("optional"):
                    raise
                values = []                  # v2-лист (Зоны) в таблице v1 — пустой справочник
            rows = read_rows(values, spec["columns"], spec["sheet"]) if values else []
            ents = {}
            for _, r in rows:
                id_ = str(r.get("id") or "").strip().lower()
                if not id_:
                    continue
                if lib == "profiles" and r.get("active") not in ("", None):
                    r["active"] = bool(parse_bool(r["active"]))
                ents[id_] = {**r, "id": id_}
            data[lib] = ents
        return Library(data, home=self.home)

    def library_issues(self, lib: Library) -> list[str]:
        """Ошибки в справочниках (ID, якоря, слоты) — для отчёта sync/check."""
        out = []
        for name, spec in SCHEMA["libraries"].items():
            pat = LIM["plate_id_re"] if name == "plates" else LIM["lib_id_re"]
            for id_, r in lib.data.get(name, {}).items():
                if not re.fullmatch(pat, id_):
                    out.append(f"{spec['sheet']}: ID «{id_}» — только латиница, цифры, _ и -"
                               + (" (вид локация/зона)" if name == "plates" else ""))
                if name == "plates":
                    try:
                        parse_anchors(r.get("anchors"))
                        parse_slots(r.get("slots"))
                    except CardError as e:
                        out.append(f"Локации, {id_}: {e}")
        return out

    def _active_profiles(self, lib: Library) -> set:
        return {i for i, r in lib.data["profiles"].items() if r.get("active") is not False}

    # --- сцены -------------------------------------------------------------------------
    def read_scenes(self) -> list[tuple[int, dict]]:
        sheet = SCHEMA["scenes"]["sheet"]
        values = self.backend.read(sheet)
        if values:
            self._cols = header_map(values[0], S_COLS, sheet)
        rows = read_rows(values, S_COLS, sheet)
        self._rowmap = {str(r.get("id") or "").strip().upper(): n for n, r in rows if r.get("id")}
        return rows

    def _do_check_files(self, override: Optional[bool]) -> bool:
        if override is not None:
            return override
        v = self.cfg.get("check_files")
        return (self.home / "refs").is_dir() if v == "auto" else bool(v)

    def check(self, check_files: Optional[bool] = None) -> tuple[Library, list[RowCheck]]:
        lib = self.read_library()
        rows = self.read_scenes()
        checks = validate(rows, lib, self._active_profiles(lib), self._do_check_files(check_files))
        return lib, checks

    # --- синхронизация -------------------------------------------------------------------
    def sync(self, db, write: bool = True, check_files: Optional[bool] = None) -> dict:
        """Таблица → БД. Возвращает отчёт {added, updated, waiting, invalid, drafts, lib_issues}."""
        lib, checks = self.check(check_files)
        rep = {"added": [], "updated": [], "waiting": [], "invalid": [], "drafts": 0,
               "skipped_changes": [], "lib_issues": self.library_issues(lib)}
        self._localize_library(lib, rep)
        lib.save(cache_path(self.home))
        statuses: dict[int, str] = {}
        current = {rc.row: str(rc.raw.get("status") or "") for rc in checks}

        for rc in checks:
            if not rc.ready:
                rep["drafts"] += 1
                if current[rc.row].startswith(STATUS["invalid"]) or current[rc.row].startswith(STATUS["waiting"]):
                    statuses[rc.row] = ""          # стёрли галочку — убрать старую ошибку
                continue
            if not rc.ok:
                rep["invalid"].append(rc.id or f"строка {rc.row}")
                statuses[rc.row] = rc.status_text()
                continue
            sid = rc.norm["id"]
            existing = db.shot(sid)
            pose_ref = self._localize_ref(rc, "pose_ref", rep)
            motion_ref = self._localize_ref(rc, "motion_ref", rep)
            location_image = self._localize_ref(rc, "location_image", rep)
            if not rc.ok:                           # ссылка не скачалась
                rep["invalid"].append(sid)
                statuses[rc.row] = rc.status_text()
                continue
            shot = to_shot(rc, pose_ref, motion_ref, location_image)
            note = ""
            if existing is None:
                sf = rc.norm.get("start_from")
                if sf:
                    aid = self._last_frame_asset(db, sf)
                    if aid is None:
                        rep["waiting"].append(sid)
                        statuses[rc.row] = STATUS["waiting"] + f"{sf} (принятое видео)"
                        continue
                    shot["end_from"] = aid
                db.add_shot(shot)
                rep["added"].append(sid)
                existing = db.shot(sid)
            else:
                old = existing.get("extra_json") or {}
                if isinstance(old, str):
                    old = json.loads(old or "{}")
                if old.get("row_hash") != shot["extra_json"]["row_hash"]:
                    note = self._apply_changes(db, existing, shot, rep)
            waiting = [d for d in (existing.get("depends_on_json") or [])
                       if (db.shot(d) or {}).get("state") != "approved"]
            text = shot_status(existing, waiting)
            if rc.warnings and existing.get("state") in ("pending", "ready"):
                text += " · ⚠ " + rc.warn_text()
            statuses[rc.row] = (text + note)[:900]

        if write:
            self._write_rows({r: {"status": s} for r, s in statuses.items() if s != current.get(r)})
        rep["statuses"] = statuses
        return rep

    def _apply_changes(self, db, existing: dict, shot: dict, rep: dict) -> str:
        sid = shot["shot_id"]
        old = existing.get("extra_json") or {}
        if isinstance(old, str):
            old = json.loads(old or "{}")
        struct_changed = (existing.get("kind") != shot["kind"] or existing.get("profile") != shot["profile"]
                          or (existing.get("depends_on_json") or []) != shot["depends_on_json"]
                          or old.get("start_from") != shot["extra_json"]["start_from"]
                          or old.get("end_with") != shot["extra_json"]["end_with"])
        if existing.get("state") in ACTIVE_STATES or (struct_changed and existing.get("state") != "pending"):
            rep["skipped_changes"].append(sid)
            return (" · ⚠ правка не применена: шот в работе" if not struct_changed
                    else " · ⚠ тип/связи/профиль после запуска не меняются — заведите новый ID")
        fields = {k: v for k, v in shot.items() if k not in ("shot_id", "end_from")}
        if existing.get("state") != "pending":     # уже идёт: структуру не трогаем (проверено выше)
            fields.pop("depends_on_json", None)
        db.update("shots", "shot_id", sid, fields)
        rep["updated"].append(sid)
        return " · правка применится к следующей попытке" if existing.get("state") != "pending" else ""

    @staticmethod
    def _last_frame_asset(db, shot_id: str) -> Optional[str]:
        """last_frame того же прогона, что принятое видео шота. Отмечается approved (см. CONTRACTS v0.4)."""
        src = db.shot(shot_id)
        if not src or src.get("state") != "approved" or not src.get("video_asset_id"):
            return None
        a = db.one("SELECT attempt_id, parent_asset_id FROM assets WHERE asset_id=?", (src["video_asset_id"],))
        while a and a.get("parent_asset_id"):           # post_fix → исходный ассет прогона
            a = db.one("SELECT attempt_id, parent_asset_id FROM assets WHERE asset_id=?", (a["parent_asset_id"],))
        if not a:
            return None
        lf = db.one("SELECT asset_id FROM assets WHERE attempt_id=? AND kind='last_frame' "
                    "ORDER BY created_at LIMIT 1", (a["attempt_id"],))
        if not lf:
            return None
        db.update("assets", "asset_id", lf["asset_id"], {"approved": 1})
        return lf["asset_id"]

    # --- файлы по ссылкам Drive ------------------------------------------------------------
    def _localize_ref(self, rc: RowCheck, key: str, rep: dict) -> Optional[str]:
        v = rc.norm.get(key)
        if not v:
            return None
        if not v.startswith("http"):
            p = _local_path(self.home, v)
            return str(p) if p else str(self.home / v)
        try:
            return str(self.download_drive(v))
        except Exception as e:        # noqa: BLE001 — сообщение пользователю
            rc.err("E_LINK_ACCESS", key, f"не скачалась ссылка: {str(e)[:160]}")
            return None

    def _localize_library(self, lib: Library, rep: dict) -> None:
        for name, spec in SCHEMA["libraries"].items():
            fkeys = [c["key"] for c in spec["columns"] if c.get("file")]
            for id_, r in lib.data.get(name, {}).items():
                for k in fkeys:
                    v = str(r.get(k) or "")
                    if v.startswith("http"):
                        try:
                            r[k] = str(self.download_drive(v))
                        except Exception as e:   # noqa: BLE001
                            rep["lib_issues"].append(f"{spec['sheet']}, {id_}: ссылка не скачалась ({str(e)[:120]})")
                            r[k] = ""

    def _session(self, scopes: list[str], prefer_user: bool = False):
        if self._creds_factory:
            return self._creds_factory(scopes, prefer_user)
        from google.auth.transport.requests import AuthorizedSession
        tok = self.home / self.drive_cfg["oauth_token"]
        if prefer_user and tok.exists():
            from google.oauth2.credentials import Credentials
            creds = Credentials.from_authorized_user_file(str(tok), scopes)
        else:
            from google.oauth2.service_account import Credentials as SACreds
            creds = SACreds.from_service_account_file(str(self.home / self.cfg["service_account"]), scopes=scopes)
        return AuthorizedSession(creds)

    def download_drive(self, url: str) -> Path:
        fid = drive_file_id(url)
        if not fid:
            raise ValueError("в ссылке нет ID файла")
        out_dir = self.home / "cache" / "refs"
        hit = sorted(out_dir.glob(f"{fid}.*")) if out_dir.exists() else []
        if hit:
            return hit[0]
        s = self._session(["https://www.googleapis.com/auth/drive.readonly"])
        meta = s.get(f"https://www.googleapis.com/drive/v3/files/{fid}",
                     params={"fields": "name,mimeType", "supportsAllDrives": "true"}, timeout=60)
        if meta.status_code in (403, 404):
            raise PermissionError("нет доступа — откройте файл для сервисного аккаунта")
        meta.raise_for_status()
        ext = Path(meta.json().get("name") or "").suffix.lower() or ".bin"
        r = s.get(f"https://www.googleapis.com/drive/v3/files/{fid}",
                  params={"alt": "media", "supportsAllDrives": "true"}, timeout=300)
        r.raise_for_status()
        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"{fid}{ext}"
        tmp = dst.with_suffix(dst.suffix + ".part")
        tmp.write_bytes(r.content)
        tmp.replace(dst)
        return dst

    # --- запись в таблицу ------------------------------------------------------------------
    def _now(self) -> str:
        try:
            from zoneinfo import ZoneInfo
            return datetime.now(ZoneInfo(self.cfg["timezone"])).strftime("%Y-%m-%d %H:%M")
        except Exception:  # noqa: BLE001
            return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    def _write_rows(self, updates: dict[int, dict]) -> None:
        if not updates:
            return
        if not self._cols:
            self.read_scenes()
        cells = []
        for row, vals in sorted(updates.items()):
            vals = {**vals, "updated": self._now()}
            for k, v in vals.items():
                if k in self._cols:
                    cells.append((row, self._cols[k] + 1, v))
        self.backend.write(SCHEMA["scenes"]["sheet"], cells)

    def on_shot_update(self, shot: dict, links: list, scores: dict) -> None:
        vals = {"status": shot_status(shot)}
        if links:
            vals["links"] = "\n".join(links)
        if scores:
            vals["scores"] = format_scores(scores)
        with self._plock:
            prev = self._pending.get(shot["shot_id"], {})
            if "links" in prev and "links" not in vals:
                vals["links"] = prev["links"]
            self._pending[shot["shot_id"]] = vals
            if self._flusher is None or not self._flusher.is_alive():
                self._flusher = threading.Thread(target=self._flush_loop, name="farm-sheets", daemon=True)
                self._flusher.start()

    def _flush_loop(self) -> None:
        while True:
            time.sleep(float(self.cfg.get("flush_s") or 3.0))
            if not self.flush():
                return

    def flush(self) -> bool:
        """Записать накопленные обновления. True — что-то было."""
        with self._plock:
            batch, self._pending = self._pending, {}
        if not batch:
            return False
        try:
            if any(sid not in self._rowmap for sid in batch):
                self.read_scenes()
            upd = {}
            for sid, vals in batch.items():
                row = self._rowmap.get(sid)
                if row is None:
                    log.warning("шот %s не найден в таблице — статус не записан", sid)
                    continue
                upd[row] = vals
            self._write_rows(upd)
        except Exception:  # noqa: BLE001 — таблица не должна ронять ферму
            log.exception("запись статусов в таблицу не удалась — повторю")
            with self._plock:
                for sid, vals in batch.items():
                    self._pending.setdefault(sid, vals)
        return True

    # --- выгрузка финала на Drive ----------------------------------------------------------
    def upload_final(self, asset_path: str, shot: dict) -> Optional[str]:
        """Финал → папка Drive пользователя (OAuth-токен, scope drive.file). Нет токена → None.

        У сервисного аккаунта нет квоты на «Мой диск» обычного Gmail (storageQuotaExceeded),
        поэтому выгрузка идёт от имени пользователя. Папку «Ферма — финалы» создаёт сама ферма
        (drive.file видит только свои файлы), id запоминается в cache/drive_folder.json.
        """
        tok = self.home / self.drive_cfg["oauth_token"]
        if not tok.exists() and not self._creds_factory:
            log.warning("нет %s — финалы на Drive не выгружаются (docs/SHEETS_SETUP.md, шаг 6)", tok)
            return None
        p = Path(asset_path)
        extra = shot.get("extra_json") or {}
        if isinstance(extra, str):
            extra = json.loads(extra or "{}")
        prefix = f"{extra.get('batch')}_" if extra.get("batch") else ""
        name = f"{prefix}{shot['shot_id']}{p.suffix}"
        s = self._session(["https://www.googleapis.com/auth/drive.file"], prefer_user=True)
        folder = self.drive_cfg.get("folder_id") or self._ensure_folder(s)
        init = s.post("https://www.googleapis.com/upload/drive/v3/files",
                      params={"uploadType": "resumable", "fields": "id,webViewLink,md5Checksum"},
                      json={"name": name, "parents": [folder]}, timeout=60)
        init.raise_for_status()
        with open(p, "rb") as f:
            r = s.put(init.headers["Location"], data=f, timeout=1800)
        r.raise_for_status()
        meta = r.json()
        md5 = hashlib.md5(p.read_bytes()).hexdigest()
        if meta.get("md5Checksum") != md5:
            s.delete(f"https://www.googleapis.com/drive/v3/files/{meta['id']}")
            raise IOError(f"md5 на Drive не совпал для {name}")
        return meta.get("webViewLink") or f"https://drive.google.com/file/d/{meta['id']}/view"

    def _ensure_folder(self, s) -> str:
        cache = self.home / "cache" / "drive_folder.json"
        if cache.exists():
            return json.loads(cache.read_text("utf-8"))["id"]
        r = s.post("https://www.googleapis.com/drive/v3/files", params={"fields": "id"},
                   json={"name": "Ферма — финалы", "mimeType": "application/vnd.google-apps.folder"}, timeout=60)
        r.raise_for_status()
        fid = r.json()["id"]
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(json.dumps({"id": fid}), "utf-8")
        return fid


# ======================================================================== модульный уровень
_inst: Optional[Sheets] = None
_inst_lock = threading.Lock()
_off_until = 0.0            # таблица не настроена — не перепроверять конфиг чаще раза в 5 минут


def make_backend(cfg: dict, home: Path):
    sc = deep_merge(DEFAULTS, cfg.get("sheets"))
    if not sc.get("enabled"):
        return None
    if sc.get("xlsx"):
        p = Path(sc["xlsx"])
        return XlsxBackend(p if p.is_absolute() else home / p)
    sa = home / sc["service_account"]
    if not sc.get("spreadsheet_id") or not sa.exists():
        return None
    return GSheetsBackend(sc["spreadsheet_id"], sa)


def get_sheets(cfg: Optional[dict] = None, home: Optional[Path] = None) -> Optional[Sheets]:
    """Экземпляр по настройкам $FARM_HOME/farm.yaml; None — таблица не настроена."""
    global _inst, _off_until
    with _inst_lock:
        if cfg is None and _inst is None and time.monotonic() < _off_until:
            return None
        if _inst is None or cfg is not None:
            home = home or farm_home()
            cfg = cfg if cfg is not None else load_config(home=home)
            try:
                be = make_backend(cfg, home)
            except Exception:  # noqa: BLE001
                log.exception("таблица недоступна")
                be = None
            if be is None:
                log.warning("таблица не настроена (sheets.spreadsheet_id + secrets/google_sa.json) — статусы не пишутся")
                _off_until = time.monotonic() + 300
                return None
            _inst = Sheets(be, cfg, home)
        return _inst


def on_shot_update(shot: dict, links: list, scores: dict) -> None:
    s = get_sheets()
    if s is not None:
        s.on_shot_update(shot, links, scores)


def upload_final(asset_path: str, shot: dict) -> Optional[str]:
    s = get_sheets()
    return s.upload_final(asset_path, shot) if s is not None else None


def sync(db=None, write: bool = True) -> Optional[dict]:
    s = get_sheets()
    if s is None:
        return None
    if db is None:
        from .db import DB
        db = DB(farm_home() / "db" / "farm.sqlite")
    return s.sync(db, write=write)


# ======================================================================== консоль
def _print_checks(lib: Library, checks: list[RowCheck], lib_issues: list[str]) -> int:
    bad = 0
    for rc in checks:
        mark = "✓" if rc.ok else "✗"
        draft = "" if rc.ready else " (черновик)"
        print(f"{mark} {rc.row:>4} {rc.id or '—'}{draft}")
        for i in rc.issues:
            print(f"       {'ошибка' if i.severity == 'error' else 'внимание'}: {i.key}: {i.msg}")
        bad += (not rc.ok) and rc.ready
    for s in lib_issues:
        print("справочник:", s)
    ok = sum(1 for r in checks if r.ok and r.ready)
    print(f"итого: в работе {sum(r.ready for r in checks)}, без ошибок {ok}, с ошибками {bad}, "
          f"черновиков {sum(not r.ready for r in checks)}")
    return 1 if bad or lib_issues else 0


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m farm.sheets", description="Таблица сцен фермы")
    ap.add_argument("cmd", choices=["check", "sync"])
    ap.add_argument("--xlsx", help="читать выгрузку .xlsx вместо Google Sheets")
    ap.add_argument("--files", action="store_true", help="проверять файлы в refs/")
    ap.add_argument("--no-write", action="store_true", help="не писать статусы в таблицу")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    home = farm_home()
    cfg = load_config(home=home)
    if a.xlsx:
        cfg = deep_merge(cfg, {"sheets": {"xlsx": a.xlsx}})
    s = get_sheets(cfg, home)
    if s is None:
        print("таблица не настроена: см. docs/SHEETS_SETUP.md")
        return 2
    if a.cmd == "check":
        lib, checks = s.check(True if a.files else None)
        return _print_checks(lib, checks, s.library_issues(lib))
    from .db import DB
    rep = s.sync(DB(home / "db" / "farm.sqlite"), write=not a.no_write)
    rep.pop("statuses", None)
    print(json.dumps(rep, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
